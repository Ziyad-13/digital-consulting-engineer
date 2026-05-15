import pytest
from unittest.mock import patch
from io import BytesIO
from openpyxl import load_workbook

from excel_export import build_boq_workbook

def test_build_boq_workbook_basic():
    project = {
        "name": "Test Project",
        "location": "Riyadh",
        "budget": 1000000,
    }
    boq_items = [
        {"category": "structural", "item_name": "Concrete", "quantity": 100, "unit": "m3"},
        {"category": "structural", "item_name": "Steel", "quantity": 50, "unit": "ton"},
        {"category": "mep", "item_name": "Pipes", "quantity": 200, "unit": "m"},
    ]

    with patch('i18n.current_lang', return_value='en'):
        result_bytes = build_boq_workbook(project, boq_items, floors=2, finish_level="luxury")

    assert isinstance(result_bytes, bytes)

    # Load workbook to verify contents
    wb = load_workbook(BytesIO(result_bytes))

    # Check sheets
    sheet_names = wb.sheetnames
    assert sheet_names[0] == "📋 Summary"
    assert "🏗️ Structural" in sheet_names
    assert "🔧 Plumbing & HVAC" in sheet_names

    # Check summary sheet
    ws_summary = wb["📋 Summary"]
    assert ws_summary.cell(row=1, column=1).value == "Bill of Quantities — Summary"
    assert ws_summary.cell(row=2, column=1).value == "Project: Test Project"

    # Check values in breakdown
    # structural is at r=16
    assert ws_summary.cell(row=16, column=1).value == "🏗️ Structural"
    assert ws_summary.cell(row=16, column=2).value == 2

    # mep is at r=17
    assert ws_summary.cell(row=17, column=1).value == "🔧 Plumbing & HVAC"
    assert ws_summary.cell(row=17, column=2).value == 1

    # total row is calculated as head_row + 1 + len(grouped) = 15 + 1 + 2 = 18
    assert ws_summary.cell(row=18, column=1).value == "TOTAL"
    assert ws_summary.cell(row=18, column=2).value == 3

    # Check category sheet
    ws_structural = wb["🏗️ Structural"]
    assert ws_structural.cell(row=1, column=1).value == "🏗️ Structural"

    # Check translation of items. The default is to return the key if not found.
    # We may need to verify the raw string if we didn't mock properly
    assert "Concrete" in str(ws_structural.cell(row=5, column=1).value) or ws_structural.cell(row=5, column=1).value == "Concrete"
    assert ws_structural.cell(row=5, column=2).value == 100
    assert ws_structural.cell(row=5, column=3).value == "m3"

def test_build_boq_workbook_empty_items():
    project = {
        "name": "Empty Project",
        "location": "Jeddah",
        "budget": 500000,
    }
    boq_items = []

    with patch('i18n.current_lang', return_value='en'):
        result_bytes = build_boq_workbook(project, boq_items)

    assert isinstance(result_bytes, bytes)
    wb = load_workbook(BytesIO(result_bytes))

    # Should only have summary sheet if no items match categories
    assert len(wb.sheetnames) == 1
    assert wb.sheetnames[0] == "📋 Summary"

def test_build_boq_workbook_rtl():
    project = {
        "name": "Arabic Project",
    }
    boq_items = [{"category": "structural", "item_name": "Concrete", "quantity": 100, "unit": "m3"}]

    with patch('i18n.current_lang', return_value='ar'):
        result_bytes = build_boq_workbook(project, boq_items)

    wb = load_workbook(BytesIO(result_bytes))

    # Check RTL is enabled on sheets
    ws_summary = wb.worksheets[0]
    assert ws_summary.sheet_view.rightToLeft is True

    ws_structural = wb.worksheets[1]
    assert ws_structural.sheet_view.rightToLeft is True
